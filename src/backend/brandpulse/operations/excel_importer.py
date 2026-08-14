"""内部经营数据 Excel 导入。

只接收工作簿中的“内部经营数据”工作表，并要求门店 ID、品牌 ID 能在
数据库主数据中匹配；未匹配的行会被拒绝，不会生成虚构门店。
"""
from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, List, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import text

from brandpulse.db_clients.postgres_client import PostgresClient

SHEET_NAME = "内部经营数据"
ALIASES = {
    "record_id": {"record_id", "op_id", "记录ID", "经营记录ID"},
    "brand_id": {"brand_id", "品牌ID"},
    "store_id": {"store_id", "门店ID"},
    "record_date": {"record_date", "日期", "经营日期"},
    "sales_amount": {"sales_amount", "销售额", "销售金额"},
    "order_count": {"order_count", "订单数", "订单量"},
    "customer_price": {"customer_price", "客单价"},
    "customer_flow": {"customer_flow", "客流量"},
    "rent": {"rent", "租金"},
    "property_fee": {"property_fee", "物业费"},
    "energy_cost": {"energy_cost", "能耗费"},
    "store_area": {"store_area", "门店面积", "面积"},
    "rent_to_sales_ratio": {"rent_to_sales_ratio", "租售比"},
    "sales_per_sqm": {"sales_per_sqm", "坪效", "每平米销售额"},
    "contract_start": {"contract_start", "合同开始日期"},
    "contract_end": {"contract_end", "contract_end_date", "合同结束日期"},
    "is_in_contract": {"is_in_contract", "合同内", "是否在合同期"},
    "data_source": {"data_source", "数据来源"},
}
NUMERIC_FIELDS = {
    "sales_amount", "customer_price", "rent", "property_fee", "energy_cost",
    "store_area", "rent_to_sales_ratio", "sales_per_sqm",
}
INTEGER_FIELDS = {"order_count", "customer_flow"}
DATE_FIELDS = {"record_date", "contract_start", "contract_end"}

TEMPLATE_HEADERS = [
    "record_id",
    "brand_id",
    "store_id",
    "record_date",
    "sales_amount",
    "order_count",
    "customer_flow",
    "store_area",
    "rent",
    "property_fee",
    "energy_cost",
    "contract_start",
    "contract_end",
    "is_in_contract",
    "data_source",
]

TEMPLATE_FIELDS = [
    ("record_id", "是", "POS/经营系统中的唯一记录 ID；重复导入时按该 ID 更新"),
    ("brand_id", "是", "必须与 brands.brand_id 完全一致"),
    ("store_id", "是", "必须与 stores.store_id 完全一致，且属于同一 brand_id"),
    ("record_date", "是", "经营日期，格式 YYYY-MM-DD"),
    ("sales_amount", "建议", "当日含税销售额；必须为非负数"),
    ("order_count", "建议", "当日订单数；用于计算真实客单价"),
    ("customer_flow", "否", "当日客流量"),
    ("store_area", "建议", "门店经营面积；用于计算坪效"),
    ("rent", "否", "与 record_date 相同统计周期口径的租金"),
    ("property_fee", "否", "物业费"),
    ("energy_cost", "否", "能耗费"),
    ("contract_start", "否", "合同开始日期，格式 YYYY-MM-DD"),
    ("contract_end", "否", "合同结束日期，格式 YYYY-MM-DD"),
    ("is_in_contract", "否", "是/否、true/false 或 1/0；留空按是处理"),
    ("data_source", "建议", "真实来源系统名称，例如 POS:系统名称；不要填写测试数据"),
]


def build_template() -> bytes:
    """生成不包含示例经营数值的空白 POS 导入模板。"""
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = SHEET_NAME
    data_sheet.append(TEMPLATE_HEADERS)
    data_sheet.freeze_panes = "A2"
    data_sheet.auto_filter.ref = f"A1:O1"
    for index, header in enumerate(TEMPLATE_HEADERS, start=1):
        data_sheet.cell(row=1, column=index).font = Font(bold=True)
        data_sheet.column_dimensions[data_sheet.cell(row=1, column=index).column_letter].width = max(14, len(header) + 3)

    note_sheet = workbook.create_sheet("字段说明")
    note_sheet.append(["字段", "是否必填", "说明"])
    for field in TEMPLATE_FIELDS:
        note_sheet.append(field)
    note_sheet.freeze_panes = "A2"
    note_sheet.column_dimensions["A"].width = 22
    note_sheet.column_dimensions["B"].width = 12
    note_sheet.column_dimensions["C"].width = 72

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _normalise_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "").replace("　", "")


def _cell_number(cell: Any, field: str) -> Any:
    value = cell.value
    if value in (None, ""):
        return None
    if isinstance(value, str):
        value = value.strip()
        # 只接受 Excel 公式中的数值常量；不在服务端执行任意公式。
        if value.startswith("="):
            value = value[1:].strip()
        if value.endswith("%"):
            return float(value[:-1])
        return float(value)
    number = float(value)
    if field == "rent_to_sales_ratio" and "%" in str(getattr(cell, "number_format", "")):
        number *= 100
    return number


def _cell_date(value: Any) -> Any:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value).strip().replace("/", "-")
    try:
        return date.fromisoformat(raw[:10])
    except ValueError as exc:
        raise ValueError(f"无法识别日期: {value}") from exc


def _find_sheet(workbook):
    if SHEET_NAME in workbook.sheetnames:
        return workbook[SHEET_NAME]
    raise ValueError(f"工作簿缺少必需工作表“{SHEET_NAME}”")


def _header_map(sheet) -> Dict[str, int]:
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10)):
        result: Dict[str, int] = {}
        for index, cell in enumerate(row, start=1):
            header = _normalise_header(cell.value)
            for field, names in ALIASES.items():
                if header in {_normalise_header(name) for name in names}:
                    result[field] = index
        if {"brand_id", "store_id", "record_date"}.issubset(result):
            return result
    raise ValueError("未找到包含 brand_id、store_id、record_date 的表头")


def parse_workbook(content: bytes) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    sheet = _find_sheet(workbook)
    headers = _header_map(sheet)
    rows: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []

    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        if not any(cell.value not in (None, "") for cell in cells):
            continue
        try:
            values = {field: cells[index - 1] for field, index in headers.items()}
            def value_of(field: str) -> Any:
                cell = values.get(field)
                return cell.value if cell is not None else None

            row: Dict[str, Any] = {
                "op_id": str(value_of("record_id") or "").strip(),
                "brand_id": str(value_of("brand_id") or "").strip(),
                "store_id": str(value_of("store_id") or "").strip(),
                "data_source": str(value_of("data_source") or "Excel:内部经营数据").strip(),
                "is_in_contract": True,
            }
            if not row["op_id"]:
                raise ValueError("record_id 不能为空；请使用业务系统中的经营记录 ID")
            if not row["brand_id"] or not row["store_id"]:
                raise ValueError("brand_id 和 store_id 不能为空")
            for field in DATE_FIELDS:
                if field in values:
                    row[field] = _cell_date(values[field].value)
            if not row.get("record_date"):
                raise ValueError("record_date 不能为空")
            for field in NUMERIC_FIELDS:
                if field in values:
                    row[field] = _cell_number(values[field], field)
            for field in INTEGER_FIELDS:
                if field in values and values[field].value not in (None, ""):
                    row[field] = int(_cell_number(values[field], field))
            if "is_in_contract" in values and value_of("is_in_contract") not in (None, ""):
                value = str(value_of("is_in_contract")).strip().lower()
                row["is_in_contract"] = value not in {"0", "false", "否", "no"}

            if row.get("customer_price") is None and row.get("sales_amount") is not None and row.get("order_count"):
                row["customer_price"] = row["sales_amount"] / row["order_count"]
            if row.get("sales_per_sqm") is None and row.get("sales_amount") is not None and row.get("store_area"):
                row["sales_per_sqm"] = row["sales_amount"] / row["store_area"]
            if row.get("rent_to_sales_ratio") is None and row.get("rent") is not None and row.get("sales_amount"):
                row["rent_to_sales_ratio"] = row["rent"] / row["sales_amount"] * 100
            for field in NUMERIC_FIELDS | INTEGER_FIELDS:
                if row.get(field) is not None and row[field] < 0:
                    raise ValueError(f"{field} 不能为负数")
            rows.append(row)
        except (KeyError, TypeError, ValueError, OverflowError) as exc:
            errors.append({"row": row_number, "error": str(exc)})

    # 主数据校验放在解析阶段，防止把不存在的门店写入经营表。
    if rows:
        client = PostgresClient()
        pairs = {(row["store_id"], row["brand_id"]) for row in rows}
        with client.engine.connect() as conn:
            known = conn.execute(text("""
                SELECT store_id, brand_id FROM stores
                WHERE store_id = ANY(CAST(:store_ids AS VARCHAR[]))
                  AND brand_id = ANY(CAST(:brand_ids AS VARCHAR[]))
            """), {
                "store_ids": list({pair[0] for pair in pairs}),
                "brand_ids": list({pair[1] for pair in pairs}),
            }).fetchall()
        known_pairs = {(str(row[0]), str(row[1])) for row in known}
        valid_rows = []
        for row in rows:
            if (row["store_id"], row["brand_id"]) not in known_pairs:
                errors.append({"row": None, "record_id": row["op_id"], "error": "store_id/brand_id 不存在或不匹配"})
            else:
                valid_rows.append(row)
        rows = valid_rows
    return rows, errors
