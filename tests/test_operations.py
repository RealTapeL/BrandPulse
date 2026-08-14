"""内部经营数据接口契约测试。"""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from brandpulse.api.app import app
from tests.conftest import AUTH_HEADERS


client = TestClient(app, headers=AUTH_HEADERS)


def test_download_pos_template_contains_headers_but_no_mock_rows():
    response = client.get("/api/v1/operations/template")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    sheet = workbook["内部经营数据"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert headers[:4] == ["record_id", "brand_id", "store_id", "record_date"]
    assert sheet.max_row == 1


def test_sales_trend_has_real_data_boundary():
    response = client.get("/api/v1/operations/metrics/sales-trend")
    assert response.status_code == 200
    assert response.json()["meta"]["source"] == "store_operations"
    assert isinstance(response.json()["series"], list)


def test_operations_readiness_uses_existing_scope():
    scopes = client.get("/api/v1/monitoring/scopes")
    assert scopes.status_code == 200
    if not scopes.json()["items"]:
        pytest.skip("数据库尚未登记监测范围")
    scope_id = scopes.json()["items"][0]["scope_id"]

    readiness = client.get("/api/v1/operations/readiness", params={"scope_id": scope_id})
    assert readiness.status_code == 200
    assert "ready_for_sales_metric" in readiness.json()


def test_sales_trend_rejects_reversed_date_range():
    response = client.get(
        "/api/v1/operations/metrics/sales-trend",
        params={"start_date": "2026-08-02", "end_date": "2026-08-01"},
    )
    assert response.status_code == 422
